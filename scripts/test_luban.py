"""Generate isolated workbooks and validate them with an actual Luban executable."""
import argparse
import json
import shutil
import subprocess
from pathlib import Path
from common import SKILL
from openpyxl import Workbook, load_workbook
from table_profile import load_profile

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--luban',type=Path,required=True)
    p.add_argument('--dotnet',default=shutil.which('dotnet'))
    p.add_argument('--out',type=Path,required=True)
    a=p.parse_args()
    if not a.dotnet:p.error('dotnet is required')
    root=a.out.resolve()
    if root.exists() and any(root.iterdir()):p.error('--out must be new or empty')
    (root/'Datas').mkdir(parents=True,exist_ok=True)
    (root/'Defines').mkdir()

    def book(name,comments,fields,types,rows,groups=None):
        wb=Workbook();ws=wb.active;ws.title=name
        for row in [['##comment',*comments],['##var',*fields],['##type',*types]]:ws.append(row)
        if groups:ws.append(['##group',*groups])
        for row in rows:ws.append([None,*row])
        ws2=wb.create_sheet('说明');ws2['A1']='辅助页，不参与导出'
        dest=root/'Datas'/f'{name}.xlsx';wb.save(dest)
        return dest

    item=book('Item',['编号','名称','编码','品质','价格','启用','引用物品','可空数量','数值列表','服务端值'],
        ['id','name','code','quality','price','enabled','friend_id','nullable_count','values#format=lite','server_value'],
        ['int!','string','string!','Quality','int#range=[0,]','bool','int#ref=TbItem?','int?','list,int','int'],
        [[101,'基础物品','01','普通',3,'true',0,None,'{1,3}',99],[102,'高级物品','02','RARE',5,'false',101,0,'{}',88]],
        ['','','','','','','','','','s'])
    skin=book('Skin',['编号','颜色','数量'],['id','color','count'],['int','string','int'],[[101,'01',2],[101,'02',1]])
    book('Names',['名称'],['name'],['string'],[['甲'],['乙']])
    book('Global',['初始货币','工资'],['initial_balance','wage'],['int','int'],[[100,20]])
    (root/'Defines/schema.xml').write_text('''<module name="">
  <enum name="Quality"><var name="NORMAL" alias="普通" value="1"/><var name="RARE" value="2"/></enum>
  <table name="TbItem" value="Item" input="Item.xlsx" index="id" readSchemaFromFile="true"/>
  <table name="TbSkin" value="Skin" input="Skin.xlsx" mode="list" index="id+color" readSchemaFromFile="true"/>
  <table name="TbNames" value="NameRow" input="Names.xlsx" mode="list" readSchemaFromFile="true"/>
  <table name="TbGlobal" value="GlobalRow" input="Global.xlsx" mode="one" readSchemaFromFile="true"/>
</module>''',encoding='utf-8')
    conf={'groups':[{'names':['c'],'default':True},{'names':['s'],'default':True}],
          'schemaFiles':[{'fileName':'Defines','type':''}],'dataDir':'Datas',
          'targets':[{'name':'client','manager':'Tables','groups':['c'],'topModule':'cfg'},
                     {'name':'server','manager':'Tables','groups':['s'],'topModule':'cfg'},
                     {'name':'all','manager':'Tables','groups':['c','s'],'topModule':'cfg'}]}
    (root/'luban.conf').write_text(json.dumps(conf,indent=2),encoding='utf-8')
    base=[a.dotnet,str(a.luban.resolve()),'--conf',str(root/'luban.conf'),'--strict','--errorFormat','json']
    checks=[]
    def invoke(name,target,generate=False):
        args=base+['-t',target]
        if generate:args+=['-c','cs-simple-json','-d','json','-x',f'outputCodeDir={root/name/"Code"}','-x',f'outputDataDir={root/name/"Data"}']
        else:args+=['-f','-x','outputSaver=null']
        r=subprocess.run(args,cwd=root,capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=60)
        (root/(name+'.log')).write_text(r.stdout+r.stderr,encoding='utf-8')
        return r
    for target in ('client','server'):
        result=invoke(target,target,True)
        assert result.returncode==0,result.stdout+result.stderr
        assert len(list((root/target/'Data').glob('*.json')))==4
        assert list((root/target/'Code').glob('*.cs'))
    def data(target,table):
        paths=[x for x in (root/target/'Data').glob('*.json') if x.stem.lower()==table.lower()]
        assert len(paths)==1,(table,paths)
        return json.loads(paths[0].read_text(encoding='utf-8-sig'))
    rows=data('client','TbItem')
    assert rows[0]['id']==101 and rows[0]['code']=='01'
    assert rows[0]['quality']==1 and rows[0]['enabled'] is True
    assert 'nullable_count' not in rows[0] and rows[1]['nullable_count']==0
    assert rows[0]['values']==[1,3] and rows[1]['values']==[]
    assert rows[1]['friend_id']==101 and 'server_value' not in rows[0]
    assert data('server','TbItem')[0]['server_value']==99
    assert len(data('client','TbSkin'))==2
    assert 'id' not in data('client','TbNames')[0]
    singleton=data('client','TbGlobal')
    assert len(singleton)==1 and singleton[0]['initial_balance']==100
    checks.append('Real client/server generation: two header-name rows, enum alias, bool, nullable, list, ref, composite key, list, singleton and group filtering')
    for name,path,cell,value,diagnostic in [
        ('bad-ref',item,'H6',9999,'9999'),
        ('bad-key',item,'B6',101,'101'),
        ('bad-range',item,'F5',-1,'price'),
        ('bad-enum',item,'E5','BAD_QUALITY','BAD_QUALITY'),
        ('bad-composite-key',skin,'C5','01','101')]:
        original=path.read_bytes()
        try:
            wb=load_workbook(path);wb.active[cell]=value;wb.save(path)
            result=invoke(name,'all')
            assert result.returncode!=0,name+' was not rejected'
            assert diagnostic.lower() in (result.stdout+result.stderr).lower(),name+' failed for an unexpected reason'
        finally:path.write_bytes(original)
        checks.append(name+' rejected by Luban')
    assert invoke('restored','all').returncode==0
    checks.append('Restored source files validate successfully')
    for name in ('luban','unreal-datatable','plain-json'):assert load_profile(name)['id']==name
    custom=load_profile('plain-json');custom['id']='project-test';custom['rules']=[str(SKILL/'references/table-adapters.md')]
    custom_file=root/'project-profile.json';custom_file.write_text(json.dumps(custom),encoding='utf-8')
    assert load_profile(str(custom_file))['id']=='project-test'
    del custom['mapping']['nullable'];custom_file.write_text(json.dumps(custom),encoding='utf-8')
    try:load_profile(str(custom_file))
    except ValueError:pass
    else:raise AssertionError('Incomplete adapter mapping was accepted')
    checks.append('Built-in and custom profiles load; incomplete mapping is rejected')
    report={'passed':True,'checks':checks,'output':str(root)}
    (root/'test-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
