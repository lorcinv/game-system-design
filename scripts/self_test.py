"""Exercise rendered output and rejection paths in an isolated directory."""
import argparse
import copy
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from common import SKILL
from build_doc import build, child
from check_docx import check
from docx import Document
from docx.shared import Pt


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--node', default=shutil.which('node'))
    p.add_argument('--out', type=Path)
    a=p.parse_args()
    if not a.node:p.error('Provide --node from the workspace runtime')
    out=a.out or Path(tempfile.mkdtemp(prefix='system-design-qa-'))
    out.mkdir(parents=True,exist_ok=True)
    passed=[]

    def draw(name,spec,expect=None):
        source=out/(name+'.json'); source.write_text(json.dumps(spec,ensure_ascii=False),encoding='utf-8')
        r=subprocess.run([a.node,str(SKILL/'scripts/draw.mjs'),str(source),str(out/name)],capture_output=True,text=True,encoding='utf-8')
        if expect:
            assert r.returncode!=0 and expect in r.stderr,(name,r.stdout,r.stderr)
        else:
            assert r.returncode==0,(name,r.stdout,r.stderr)
            data=json.loads((out/(name+'.layout.json')).read_text(encoding='utf-8'))
            assert data['minFontPt']>=12 and (out/(name+'.png')).read_bytes().startswith(b'\x89PNG'),name
            return data

    ui=json.loads((SKILL/'evals/ui.json').read_text(encoding='utf-8'))
    flow=json.loads((SKILL/'evals/flow.json').read_text(encoding='utf-8'))
    data=draw('ui',ui)
    assert set(data['callouts'])=={'1','2','3'}
    passed.append('ui layout, font size and callouts')
    first=hashlib.sha256((out/'ui.svg').read_bytes()).hexdigest()
    draw('ui-repeat',ui)
    assert hashlib.sha256((out/'ui-repeat.svg').read_bytes()).hexdigest()==first
    passed.append('repeat rendering is deterministic')
    draw('flow',flow)
    passed.append('branched flow and labels')
    long={'kind':'ui','width':600,'root':{'type':'panel','title':'长文本','children':[{'type':'text','text':'当列表中包含较长的中文物品名称时，保持完整显示并按当前区域宽度换行。'*5}]}}
    data=draw('long-chinese',long)
    assert data['height']>300
    passed.append('long Chinese text expands height')
    bad=copy.deepcopy(long);bad['height']=100
    draw('overflow',bad,'exceeds')
    passed.append('fixed canvas rejects overflow')
    bad=copy.deepcopy(ui);bad['root']['children'][1]['children'][1]['callout']=1
    draw('duplicate-callout',bad,'Duplicate callout')
    bad=copy.deepcopy(flow);bad['edges'][0]['to']='absent'
    draw('missing-node',bad,'Unknown edge endpoint')
    bad=copy.deepcopy(ui);bad['width']=1800
    draw('print-too-small',bad,'Text would be')
    passed.append('duplicate callouts, missing nodes and small print rejected')
    spec=json.loads((SKILL/'evals/document.json').read_text(encoding='utf-8'))
    doc_path=out/'document.docx';build(spec,doc_path,out)
    assert not check(doc_path)
    passed.append('document fields, styles, tables and figures')
    field_table=Document(doc_path).tables[-1]
    assert len(field_table.columns)==2
    assert [c.text for c in field_table.rows[0].cells]==['中文字段','英文字段']
    assert [c.text for c in field_table.rows[1].cells]==['物品编号','id']
    malformed={'title':'字段表检查','blocks':[{'type':'field_table','rows':[['编号','id','多余列']]}]}
    try:build(malformed,out/'bad-field-table.docx',out)
    except ValueError as e:assert 'exactly two columns' in str(e)
    else:raise AssertionError('Third field column was accepted')
    passed.append('field dictionary keeps name pairs and rejects extra columns')
    for name,patch,expected in [
        ('inherited-bullet',lambda d:child(child(d.styles['A正文-1标题']._element.get_or_add_pPr(),'numPr'),'numId',val=1),'Heading has list numbering'),
        ('paragraph-marker',lambda d:setattr(d.styles['A正文-1标题'].paragraph_format,'keep_with_next',True),'Heading has paragraph-control marker'),
        ('small-table',lambda d:setattr(d.tables[0].cell(0,0).paragraphs[0].runs[0].font,'size',Pt(9)),'Small or unresolved font size'),
        ('literal-bullet',lambda d:setattr(next(p for p in d.paragraphs if p.style.name=='A正文-1标题'),'text','● 一、概述'),'Heading has bullet character'),
    ]:
        d=Document(doc_path);patch(d);dest=out/(name+'.docx');d.save(dest)
        assert any(expected in e for e in check(dest)),(name,check(dest))
    bad=copy.deepcopy(spec);bad['blocks'].append({'type':'paragraph','runs':[{'ref':'absent'}]})
    try:build(bad,out/'bad-ref.docx',out)
    except ValueError as e:assert 'Unknown reference' in str(e)
    else:raise AssertionError('Missing reference was accepted')
    passed.append('heading bullets, small table fonts and broken refs rejected')
    long_table={'title':'跨页表格检查','blocks':[{'type':'heading','level':1,'id':'table','text':'一、字段说明'},
      {'type':'field_table','rows':[[f'参数{i}',f'value_{i}'] for i in range(1,46)]}]}
    build(long_table,out/'long-table.docx',out)
    report={'passed':True,'checks':passed,'output':str(out)}
    (out/'test-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
