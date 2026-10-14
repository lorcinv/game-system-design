"""Build a new system specification from the user's extracted Word template."""
import argparse
import json
import re
from pathlib import Path
from common import SKILL
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def child(parent, tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def field(p, instruction, cached='', link=False):
    r = p.add_run()
    child(r._r, 'fldChar', fldCharType='begin')
    code_run = p.add_run()
    if link:
        code_run.font.size = Pt(12)
        child(code_run._r.get_or_add_rPr(), 'color', val='0563C1')
        child(code_run._r.get_or_add_rPr(), 'u', val='single')
        instruction += ' \\* CHARFORMAT'
    code = child(code_run._r, 'instrText')
    code.set(qn('xml:space'), 'preserve')
    code.text = ' ' + instruction + ' '
    child(p.add_run()._r, 'fldChar', fldCharType='separate')
    r = p.add_run(cached)
    if link:
        child(r._r.get_or_add_rPr(), 'color', val='0563C1')
        child(r._r.get_or_add_rPr(), 'u', val='single')
    child(p.add_run()._r, 'fldChar', fldCharType='end')


def build(spec, output, base):
    doc = Document(SKILL / 'assets/word-template.docx')
    body = doc._element.body
    for node in list(body):
        if node.tag != qn('w:sectPr'):
            body.remove(node)
    headings = {}
    for b in spec['blocks']:
        if b['type'] == 'heading':
            key = b.get('id')
            if not key or not re.fullmatch('[A-Za-z][A-Za-z0-9_]{0,35}', key):
                raise ValueError('Heading id must be 1-36 ASCII letters/digits/underscores, starting with a letter')
            if key in headings:
                raise ValueError(f'Duplicate heading id: {key}')
            if b['level'] not in (1, 2, 3):
                raise ValueError('Heading level must be 1, 2, or 3')
            headings[key] = b['text']
    title = str(spec['title']).strip()
    if not title:
        raise ValueError('Document title is empty')
    doc.core_properties.title = title
    doc.core_properties.author = spec.get('author', '')
    doc.core_properties.last_modified_by = spec.get('author', '')
    doc.add_paragraph(title, style='A正文标题')

    def write_runs(p, items):
        for item in items:
            if isinstance(item, str):
                p.add_run(item)
            elif 'ref' in item:
                if item['ref'] not in headings:
                    raise ValueError(f"Unknown reference: {item['ref']}")
                field(p, 'REF s_' + item['ref'] + ' \\h', headings[item['ref']], link=True)
            elif 'text' in item:
                r = p.add_run(item['text'])
                r.bold = bool(item.get('bold'))
                if item.get('highlight'):
                    child(r._r.get_or_add_rPr(), 'highlight', val='yellow')
            else:
                raise ValueError(f'Invalid paragraph run: {item}')

    def table(headers, rows, widths=None):
        if not headers or any(len(row) != len(headers) for row in rows):
            raise ValueError('Table row width mismatch')
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        pr = t._tbl.tblPr
        old = pr.find(qn('w:tblBorders'))
        if old is not None:
            pr.remove(old)
        borders = child(pr, 'tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            child(borders, edge, val='single', sz=5, color='B7B7B7')
        available = (doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin)
        ratios = widths or [1] * len(headers)
        if len(ratios) != len(headers) or any(v <= 0 for v in ratios):
            raise ValueError('Invalid table column widths')
        for column, ratio in zip(t.columns, ratios):
            column.width = int(available * ratio / sum(ratios))
        for i, values in enumerate([headers] + rows):
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr()
            child(trpr, 'cantSplit')
            if i == 0:
                child(trpr, 'tblHeader')
            for j, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = int(available * ratios[j] / sum(ratios))
                p = cell.paragraphs[0]
                p.style = 'A表格'
                if isinstance(value, dict) and 'runs' in value:
                    write_runs(p, value['runs'])
                else:
                    p.add_run(str(value))
                if i == 0:
                    for r in p.runs:
                        r.bold = True
                if i == 0:
                    child(cell._tc.get_or_add_tcPr(), 'shd', fill='E7E6E6', val='clear')
        doc.add_paragraph('', style='A正文')
        return t

    if spec.get('revision'):
        table(['项目', '编辑人', '更新时间', '更新内容'], spec['revision'], [1, 1, 1.3, 2.7])
    if spec.get('toc'):
        p = doc.add_paragraph(style='A目录')
        field(p, 'TOC \\o "1-3" \\h \\z \\u')
        doc.add_page_break()

    bookmark = 1
    for b in spec['blocks']:
        kind = b['type']
        if kind == 'heading':
            if re.match(r'^\s*[•●▪◦■]', b['text']):
                raise ValueError('Heading cannot start with a bullet')
            p = doc.add_paragraph(style=f"A正文-{b['level']}标题")
            child(p._p, 'bookmarkStart', id=bookmark, name='s_' + b['id'])
            p.add_run(b['text'])
            child(p._p, 'bookmarkEnd', id=bookmark)
            bookmark += 1
        elif kind == 'paragraph':
            p = doc.add_paragraph(style='A正文')
            write_runs(p, b.get('runs', [b.get('text', '')]))
        elif kind == 'table':
            table(b['headers'], b['rows'], b.get('widths'))
        elif kind == 'field_table':
            if any(len(row) != 2 for row in b['rows']):
                raise ValueError('Field dictionary requires exactly two columns: Chinese and English')
            table(['中文字段', '英文字段'], b['rows'], b.get('widths', [1, 1]))
        elif kind == 'figure':
            img = (base / b['file']).resolve()
            layout_path = img.with_suffix('.layout.json')
            if not layout_path.is_file():
                raise ValueError(f'Figure needs renderer layout report: {layout_path}')
            layout = json.loads(layout_path.read_text(encoding='utf-8-sig'))
            width_cm = float(b.get('width_cm', 17))
            max_width = (doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin) / 360000
            points = layout['minFontPx'] * width_cm / 2.54 * 72 / layout['width']
            if width_cm <= 0 or width_cm > max_width + 0.01:
                raise ValueError('Figure exceeds document text width')
            if points < 12 - 0.001:
                raise ValueError(f'Figure text would be {points:.1f} pt; split the diagram or enlarge its text')
            if layout['height'] * width_cm / layout['width'] > 23:
                raise ValueError('Figure too tall for page; split it instead of shrinking text')
            p = doc.add_paragraph(style='A图注')
            p.paragraph_format.keep_with_next = bool(b.get('caption'))
            p.add_run().add_picture(str(img), width=Cm(width_cm))
            if b.get('caption'):
                doc.add_paragraph(b['caption'], style='A图注')
        elif kind == 'pagebreak':
            doc.add_page_break()
        else:
            raise ValueError(f'Unknown block: {kind}')
    settings = doc.settings.element
    for old in settings.findall(qn('w:updateFields')):
        settings.remove(old)
    child(settings, 'updateFields', val='true')
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    from check_docx import check
    errors = check(output)
    if errors:
        output.unlink()
        raise ValueError('; '.join(errors))
    return output


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('spec', type=Path)
    p.add_argument('output', type=Path)
    args = p.parse_args()
    build(json.loads(args.spec.read_text(encoding='utf-8-sig')), args.output, args.spec.parent)
    print(args.output)


if __name__ == '__main__':
    main()
